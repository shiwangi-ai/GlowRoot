#!/usr/bin/env python3
"""
Builds docs/GlowRoot-Project-Walkthrough.pptx — a 23-slide cross-functional
requirements walkthrough for GlowRoot.

Content is drawn from the real repo: docs/frontend.md, backend.md,
backend-auth-plan.md, customer-journey.md, user-journey.md and the source.
Palette is the app's own design tokens.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

OUT = "/Users/shiwangi/repos/GlowRoot-final/docs/GlowRoot-Project-Walkthrough.pptx"

# ---------------------------------------------------------------- design tokens
LEAF      = RGBColor(0x2F, 0x5D, 0x34)
LEAF_DEEP = RGBColor(0x19, 0x3D, 0x20)
LEAF_SOFT = RGBColor(0x5A, 0x8A, 0x4A)
CARROT    = RGBColor(0xD9, 0x77, 0x2E)
BERRY     = RGBColor(0x7A, 0x4A, 0x8F)
APPLE     = RGBColor(0xD4, 0x53, 0x3A)
GOLD      = RGBColor(0xE8, 0xB2, 0x3F)
CREAM     = RGBColor(0xF6, 0xF0, 0xE2)
CREAM2    = RGBColor(0xEF, 0xE6, 0xD2)
PAPER     = RGBColor(0xFB, 0xF8, 0xF0)
INK       = RGBColor(0x24, 0x30, 0x18)
INK_SOFT  = RGBColor(0x5B, 0x64, 0x50)
LINE      = RGBColor(0xE0, 0xD6, 0xBF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
TEAL      = RGBColor(0x4A, 0x7C, 0x8F)
DEEPCARD  = RGBColor(0x14, 0x33, 0x1B)

H_FONT = "Trebuchet MS"
B_FONT = "Calibri"
M_FONT = "Consolas"

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.72)
CW = W - 2 * MX

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------- primitives
def E(v):
    """Coerce a computed length to integer EMU. Emu arithmetic yields floats and
    a fractional EMU serialises to invalid XML."""
    return Emu(int(round(float(v))))


def slide(bgcolor=PAPER):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid(); r.fill.fore_color.rgb = bgcolor
    r.line.fill.background(); r.shadow.inherit = False
    return s


def box(s, x, y, w, h, fill=None, line=None, radius=None,
        shape=MSO_SHAPE.RECTANGLE, line_w=1.0):
    sh = s.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if sh.has_text_frame:
        sh.text_frame.text = ""
    return sh


def txt(s, x, y, w, h, runs, size=14, color=INK, bold=False, font=B_FONT,
        align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP, space_after=0,
        italic=False):
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        para.space_after = Pt(space_after)
        chunks = p if isinstance(p, list) else [(p, {})]
        for text, o in chunks:
            r = para.add_run(); r.text = text
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return tb


def header(s, kicker, title, accent=CARROT, dark=False):
    tcol = CREAM if dark else INK
    box(s, MX, Inches(0.52), Inches(0.055), Inches(0.52), fill=accent)
    txt(s, MX + Inches(0.20), Inches(0.50), CW, Inches(0.24),
        kicker.upper(), size=10.5, color=accent, bold=True, font=H_FONT)
    txt(s, MX + Inches(0.20), Inches(0.74), CW, Inches(0.42),
        title, size=24, color=tcol, bold=True, font=H_FONT, spacing=0.95)
    return Inches(1.42)


def footer(s, n, dark=False):
    col = RGBColor(0x8A, 0x93, 0x80) if not dark else RGBColor(0x7E, 0x8A, 0x74)
    txt(s, MX, H - Inches(0.50), CW - Inches(0.6), Inches(0.24),
        "GlowRoot  ·  Project walkthrough", size=9, color=col, font=B_FONT)
    txt(s, W - MX - Inches(0.6), H - Inches(0.50), Inches(0.6), Inches(0.24),
        str(n), size=9, color=col, font=B_FONT, align=PP_ALIGN.RIGHT)


def chip(s, x, y, label, color, textcolor=WHITE, w=None, size=9.5):
    w = E(w or (Inches(0.125) * len(label) + Inches(0.26)))
    sh = box(s, x, y, w, Inches(0.27), fill=color,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = True; r.font.name = H_FONT
    r.font.color.rgb = textcolor
    return E(x + w)


def bullets(s, x, y, w, items, size=13.5, gap=0.30, dot=CARROT, lead_color=INK,
            body_color=INK_SOFT):
    """One textbox, one paragraph per item, with a hanging indent so wrapped
    lines align under the text rather than under the bullet. Letting PowerPoint
    own the wrapping is what keeps the gaps between items uniform."""
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(Inches(0.4)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    hang = Inches(0.19)
    for i, it in enumerate(items):
        lead, body = (it, "") if isinstance(it, str) else it
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.line_spacing = 1.14
        para.space_after = Pt(gap * 24)
        pPr = para._p.get_or_add_pPr()
        pPr.set("marL", str(int(hang)))
        pPr.set("indent", str(-int(hang)))
        r = para.add_run(); r.text = "\u25cf  "
        r.font.name = B_FONT; r.font.size = Pt(size * 0.55)
        r.font.color.rgb = dot; r.font.bold = True
        if lead:
            r = para.add_run(); r.text = lead
            r.font.name = B_FONT; r.font.size = Pt(size)
            r.font.bold = True; r.font.color.rgb = lead_color
        if body:
            r = para.add_run(); r.text = ("  " + body) if lead else body
            r.font.name = B_FONT; r.font.size = Pt(size)
            r.font.color.rgb = body_color
    return tb


def table(s, x, y, w, rows, col_w, header_fill=LEAF, header_color=CREAM,
          size=11, header_size=10.5, row_h=0.34, head_h=0.34, zebra=CREAM,
          body_color=INK, bold_first=True):
    nrows, ncols = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nrows, ncols, E(x), E(y), E(w),
                              E(Inches(head_h + row_h * (nrows - 1))))
    tbl = gtbl.table
    tbl.first_row = True
    tbl.horz_banding = False
    total = sum(col_w)
    for i, cwt in enumerate(col_w):
        tbl.columns[i].width = E(w * (cwt / total))
    tbl.rows[0].height = Inches(head_h)
    for r in range(1, nrows):
        tbl.rows[r].height = Inches(row_h)

    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.10); cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = (header_fill if r == 0
                                        else (zebra if r % 2 == 1 else WHITE))
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            chunks = val if isinstance(val, list) else [(str(val), {})]
            for text, o in chunks:
                run = p.add_run(); run.text = text
                f = run.font
                f.name = o.get("font", B_FONT)
                f.size = Pt(o.get("size", header_size if r == 0 else size))
                f.bold = o.get("bold", r == 0 or (bold_first and c == 0))
                f.italic = o.get("italic", False)
                f.color.rgb = o.get("color", header_color if r == 0 else body_color)
    return gtbl


def mono(s, x, y, w, h, text, size=10.5, color=LEAF_DEEP, fill=WHITE, line=LINE):
    box(s, x, y, w, h, fill=fill, line=line)
    tb = s.shapes.add_textbox(E(x + Inches(0.16)), E(y + Inches(0.12)),
                              E(w - Inches(0.32)), E(h - Inches(0.24)))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.04
        r = p.add_run(); r.text = ln if ln else " "
        r.font.name = M_FONT; r.font.size = Pt(size); r.font.color.rgb = color
    return tb


def card(s, x, y, w, h, title, body, accent=LEAF, fill=WHITE, title_size=13,
         body_size=11.5, kicker=None):
    box(s, x, y, w, h, fill=fill, line=LINE)
    box(s, x, y, w, Inches(0.05), fill=accent)
    cy = y + Inches(0.20)
    if kicker:
        txt(s, x + Inches(0.20), cy, w - Inches(0.4), Inches(0.2),
            kicker.upper(), size=9, color=accent, bold=True, font=H_FONT)
        cy = E(cy + Inches(0.22))
    txt(s, x + Inches(0.20), cy, w - Inches(0.4), Inches(0.3), title,
        size=title_size, color=INK, bold=True, font=H_FONT, spacing=1.0)
    txt(s, x + Inches(0.20), E(cy + Inches(0.30)), w - Inches(0.4),
        E(h - (cy - y) - Inches(0.4)), body, size=body_size, color=INK_SOFT,
        spacing=1.16)


def status_row(state):
    cfg = {"done": ("■ Built", LEAF_SOFT),
           "partial": ("■ Partial", GOLD),
           "missing": ("■ Not built", APPLE)}[state]
    return [(cfg[0], {"color": cfg[1], "bold": True})]


N = 0
def num():
    global N
    N += 1
    return N


HALF = E((CW - Inches(0.32)) / 2)
X2 = E(MX + HALF + Inches(0.32))

# =============================================================== 1 · TITLE
s = slide(LEAF_DEEP)
box(s, 0, 0, W, Inches(0.10), fill=CARROT)
for cx, cy, d, col in [(Inches(10.3), Inches(0.9), Inches(3.1), LEAF),
                       (Inches(11.5), Inches(3.5), Inches(4.2), DEEPCARD),
                       (Inches(9.6), Inches(5.3), Inches(2.2), LEAF)]:
    box(s, cx, cy, d, d, fill=col, shape=MSO_SHAPE.OVAL)

txt(s, MX, Inches(1.55), Inches(8.6), Inches(0.3),
    "REQUIREMENTS & ARCHITECTURE WALKTHROUGH", size=11.5, color=GOLD,
    bold=True, font=H_FONT)
txt(s, MX, Inches(2.02), Inches(8.8), Inches(1.2),
    "GlowRoot", size=64, color=CREAM, bold=True, font=H_FONT, spacing=0.9)
txt(s, MX, Inches(3.24), Inches(8.3), Inches(1.0),
    "A cold-pressed juice brand, built as two products over one shared domain: "
    "a customer storefront, and the ops console that fulfils it.",
    size=16.5, color=RGBColor(0xC9, 0xD6, 0xC0), spacing=1.24)
box(s, MX, Inches(4.34), Inches(1.5), Inches(0.035), fill=CARROT)
txt(s, MX, Inches(4.62), Inches(8.9), Inches(0.9),
    [[("One deck, five audiences.  ", {"bold": True, "color": CREAM}),
      ("Architecture, UX, marketing, QA and DevOps each get a section — "
       "and one concrete ask at the end.",
       {"color": RGBColor(0xA9, 0xBA, 0xA0)})]],
    size=13.5, spacing=1.2)

x = MX
for lbl, col in [("Product", CARROT), ("Architecture", LEAF_SOFT),
                 ("UI / UX", BERRY), ("Marketing", GOLD), ("QA", APPLE),
                 ("DevOps", TEAL)]:
    x = E(chip(s, x, Inches(5.55), lbl, col,
               textcolor=INK if col is GOLD else WHITE) + Inches(0.11))

txt(s, MX, Inches(6.45), Inches(7.0), Inches(0.3),
    "Turborepo + pnpm  ·  Next.js 14  ·  TypeScript  ·  storefront :3000",
    size=10, color=RGBColor(0x86, 0x9A, 0x7C), font=M_FONT)
num()

# =============================================================== 2 · READER MAP
s = slide()
y = header(s, "How to read this deck", "Every role has a section — here is yours")
txt(s, MX, y, CW, Inches(0.5),
    "Slides 3–11 are shared context; everyone should be in the room for those. "
    "From slide 12 the deck fans out by discipline, then converges again on "
    "status, roadmap and asks.",
    size=13, color=INK_SOFT, spacing=1.2)
rows = [["Slides", "Section", "Who it's for", "What you get out of it"],
        ["3 – 6", "Problem, users, principles, scope", "Everyone",
         "Why this product exists, and what v1 deliberately excludes"],
        ["7 – 11", "Journey and the three key decisions", "Everyone",
         "The product logic the whole build hangs off"],
        ["12 – 14", "System shape, domain model, identity",
         "Senior architects, security",
         "The seams, the invariants, and the modelling debt"],
        ["15 – 16", "Design system and interaction contracts", "UI / UX designers",
         "Token set, component inventory, and the states we still owe"],
        ["17", "Positioning and funnel", "Marketing",
         "What the site does for acquisition, and the one known leak"],
        ["18", "Test strategy", "QA",
         "Critical paths, the status matrix, and what's mocked"],
        ["19", "Environments and Phase 4 runway", "DevOps",
         "What runs today, and what production actually needs"],
        ["20 – 23", "Status, roadmap, risks, asks", "Everyone",
         "Where we are, what's next, and what I need from you"]]
table(s, MX, E(y + Inches(0.72)), CW, rows, [0.85, 3.1, 2.4, 4.7],
      row_h=0.44, head_h=0.34, size=11)
footer(s, num())

# =============================================================== 3 · PROBLEM
s = slide()
y = header(s, "The problem", "Perishable goods break the e-commerce default")
cw = E((CW - Inches(0.44)) / 3)
card(s, MX, y, cw, Inches(2.55), "Juice is pressed to order",
     "There is no warehouse buffer and no shipping window to hide behind — a "
     "bottle is made the night before it's drunk. So the delivery date isn't a "
     "fulfilment detail, it's part of the purchase, and it has to be settled "
     "before payment.",
     accent=CARROT, kicker="Constraint 01")
card(s, E(MX + cw + Inches(0.22)), y, cw, Inches(2.55),
     "Two revenue modes, one funnel",
     "One-time buyers (trial, gifting) and subscribers have different "
     "economics, objections and journeys — but they must share one catalogue, "
     "one cart and one checkout, or we maintain the same flow twice.",
     accent=BERRY, kicker="Constraint 02")
card(s, E(MX + 2 * (cw + Inches(0.22))), y, cw, Inches(2.55),
     "Ops needs tomorrow's answer",
     "\"What do we press tonight, and what produce do we buy?\" A storefront "
     "can't answer that. Somebody was going to run it on a spreadsheet — which "
     "is exactly how order history and stock levels drift apart.",
     accent=LEAF, kicker="Constraint 03")

box(s, MX, E(y + Inches(2.82)), CW, Inches(1.30), fill=CREAM, line=LINE)
box(s, MX, E(y + Inches(2.82)), Inches(0.05), Inches(1.30), fill=LEAF)
txt(s, E(MX + Inches(0.28)), E(y + Inches(3.00)), Inches(2.4), Inches(0.3),
    "SO THE SHAPE IS", size=10, color=LEAF, bold=True, font=H_FONT)
txt(s, E(MX + Inches(0.28)), E(y + Inches(3.26)), E(CW - Inches(0.6)),
    Inches(0.8),
    [[("Two surfaces over one shared domain. ", {"bold": True, "color": INK}),
      ("A storefront that sells and an ops console that fulfils, reading the "
       "same products, the same orders and the same produce — so a sale on one "
       "side is visibly the same object as a press list on the other. That is "
       "the entire reason this is a monorepo and not two projects.",
       {"color": INK_SOFT})]],
    size=12.5, spacing=1.2)
footer(s, num())

# =============================================================== 4 · USERS
s = slide()
y = header(s, "Who we're building for",
           "Three users, three different definitions of success", accent=BERRY)
cw = E((CW - Inches(0.44)) / 3)
personas = [
    ("Aditi · the gifter", "ONE-TIME", CARROT,
     "Wants to try it, or send a box to someone. Zero appetite for commitment, "
     "and none for creating an account.",
     "A frictionless first order. She should never meet a form field whose "
     "purpose she can't guess.",
     "Combos, gift boxes, build-your-own box,\nguest-friendly checkout"),
    ("Rohan · the subscriber", "RECURRING", BERRY,
     "Drinks one most mornings. Wants it handled, but on his terms — his days, "
     "his cadence, pausable when he travels.",
     "Set and forget, with control. He should be able to change his mind "
     "without emailing anyone.",
     "Plans, weekly/monthly cadence, start date\n+ delivery days, account area"),
    ("Meera · ops lead", "INTERNAL", LEAF,
     "Runs the press and the produce order. Today she'd be reconciling a "
     "spreadsheet against WhatsApp messages.",
     "One screen that tells her what to press, what's moving, and what's "
     "running low.",
     "Orders list, status changes, KPI dashboard,\nmaterials + low stock"),
]
for i, (name, tag, col, ctx, success, surface) in enumerate(personas):
    x = E(MX + i * (cw + Inches(0.22)))
    box(s, x, y, cw, Inches(4.00), fill=WHITE, line=LINE)
    box(s, x, y, cw, Inches(0.05), fill=col)
    chip(s, E(x + Inches(0.20)), E(y + Inches(0.22)), tag, col, size=8.5)
    txt(s, E(x + Inches(0.20)), E(y + Inches(0.62)), E(cw - Inches(0.4)),
        Inches(0.32), name, size=15, color=INK, bold=True, font=H_FONT)
    txt(s, E(x + Inches(0.20)), E(y + Inches(1.00)), E(cw - Inches(0.4)),
        Inches(0.9), ctx, size=11.5, color=INK_SOFT, spacing=1.16)
    box(s, E(x + Inches(0.20)), E(y + Inches(1.94)), E(cw - Inches(0.4)),
        Inches(0.012), fill=LINE)
    txt(s, E(x + Inches(0.20)), E(y + Inches(2.08)), E(cw - Inches(0.4)),
        Inches(0.2), "SUCCESS LOOKS LIKE", size=8.5, color=col, bold=True,
        font=H_FONT)
    txt(s, E(x + Inches(0.20)), E(y + Inches(2.30)), E(cw - Inches(0.4)),
        Inches(0.85), success, size=11.5, color=INK, spacing=1.16)
    txt(s, E(x + Inches(0.20)), E(y + Inches(3.26)), E(cw - Inches(0.4)),
        Inches(0.2), "SURFACES THEY TOUCH", size=8.5, color=col, bold=True,
        font=H_FONT)
    txt(s, E(x + Inches(0.20)), E(y + Inches(3.48)), E(cw - Inches(0.4)),
        Inches(0.45), surface, size=9.5, color=INK_SOFT, spacing=1.16,
        font=M_FONT)
footer(s, num())

# =============================================================== 5 · PRINCIPLES
s = slide()
y = header(s, "Product principles", "Four rules that settle most arguments for us")
txt(s, MX, y, CW, Inches(0.36),
    "These are decisions, not aspirations. Each one is already enforced "
    "somewhere in the code — and each one has a cost we took knowingly.",
    size=12.5, color=INK_SOFT, spacing=1.2)
items = [
    ("01", "Never gate the catalogue.",
     "No pincode wall and no login wall in front of browsing. The landing page "
     "is the strongest marketing asset we have; gating it costs first "
     "impressions and search visibility.",
     "Cost: people out of our area can browse things they can't buy — so the "
     "soft banner has to do real work."),
    ("02", "Confirm serviceability before money, not before browsing.",
     "One hard check, at checkout, inside the server action. Everywhere else, "
     "serviceability is only a hint.",
     "Cost: a customer can fill an entire checkout form before being told no. "
     "Accepted — see slide 9."),
    ("03", "Delivery date is a product decision.",
     "Date and slot are a first-class, numbered checkout step — not a shipping "
     "estimate tucked onto the confirmation page.",
     "Cost: checkout is longer than a standard cart. Worth it for something "
     "made the night before."),
    ("04", "Every sale is traceable to the produce it consumes.",
     "Each product carries a recipe: a bill of materials against the raw "
     "produce ops actually tracks.",
     "Cost: the catalogue is more work to author. It's what makes the ops side "
     "more than a list of orders."),
]
cy = E(y + Inches(0.56))
for n, title, body, cost in items:
    box(s, MX, cy, CW, Inches(1.12), fill=WHITE, line=LINE)
    box(s, MX, cy, Inches(0.05), Inches(1.12), fill=LEAF)
    txt(s, E(MX + Inches(0.24)), E(cy + Inches(0.24)), Inches(0.7),
        Inches(0.5), n, size=22, color=CREAM2, bold=True, font=H_FONT)
    txt(s, E(MX + Inches(0.92)), E(cy + Inches(0.16)), E(CW - Inches(1.2)),
        Inches(0.3), title, size=14, color=INK, bold=True, font=H_FONT)
    txt(s, E(MX + Inches(0.92)), E(cy + Inches(0.45)), E(CW - Inches(1.2)),
        Inches(0.3), body, size=11.5, color=INK_SOFT, spacing=1.14)
    txt(s, E(MX + Inches(0.92)), E(cy + Inches(0.82)), E(CW - Inches(1.2)),
        Inches(0.24), cost, size=10.5, color=CARROT, spacing=1.1, italic=True)
    cy = E(cy + Inches(1.22))
footer(s, num())

# =============================================================== 6 · SCOPE
s = slide()
y = header(s, "Scope", "What v1 is — and what it deliberately is not")
CH = Inches(3.62)
box(s, MX, y, HALF, CH, fill=WHITE, line=LINE)
box(s, MX, y, HALF, Inches(0.05), fill=LEAF_SOFT)
txt(s, E(MX + Inches(0.22)), E(y + Inches(0.20)), E(HALF - Inches(0.44)),
    Inches(0.3), "IN SCOPE — SHIPPED OR IN FLIGHT", size=10, color=LEAF,
    bold=True, font=H_FONT)
bullets(s, E(MX + Inches(0.22)), E(y + Inches(0.52)), E(HALF - Inches(0.44)), [
    ("Catalogue", "15 SKUs — 13 juices, 1 smoothie, 1 shot — plus seasonal produce and detail pages"),
    ("Combos & boxes", "curated combos, and a build-your-own box that keeps different fills as separate cart lines"),
    ("Cart", "drawer, quantities, survives reload"),
    ("One-time checkout", "pincode → date + slot → address → pay, with inline validation"),
    ("Subscriptions", "3 plans, weekly/monthly cadence, start date and delivery days"),
    ("Accounts", "real passwords, server session, order history, order timeline"),
    ("Public tracking", "order ID plus email, no login required"),
    ("Ops console", "orders list and detail, status changes, KPI dashboard, materials with low-stock alerts"),
], size=11, gap=0.26, dot=LEAF_SOFT)

box(s, X2, y, HALF, CH, fill=WHITE, line=LINE)
box(s, X2, y, HALF, Inches(0.05), fill=APPLE)
txt(s, E(X2 + Inches(0.22)), E(y + Inches(0.20)), E(HALF - Inches(0.44)),
    Inches(0.3), "OUT OF SCOPE FOR V1 — NAMED, NOT FORGOTTEN", size=10,
    color=APPLE, bold=True, font=H_FONT)
bullets(s, E(X2 + Inches(0.22)), E(y + Inches(0.52)), E(HALF - Inches(0.44)), [
    ("Recurring generation", "a subscription writes one order today; nothing generates delivery two"),
    ("Inventory deduction", "recipes exist and are displayed, but stock never moves"),
    ("Admin authentication", "anyone who can reach the admin port can change any status"),
    ("Real payment gateway", "mocked deliberately, behind a swappable function"),
    ("Notifications", "no email, SMS or WhatsApp — confirmation is on-screen only"),
    ("Waitlist capture", "the out-of-area promise has nowhere to write"),
    ("Customer-side cancel", "cancelled exists as a status; only ops can set it"),
    ("Catalogue editing", "admin product and material pages are read-only views"),
], size=11, gap=0.26, dot=APPLE)

txt(s, MX, E(y + CH + Inches(0.16)), CW, Inches(0.3),
    [[("Why say the out-of-scope list out loud:  ", {"bold": True, "color": INK}),
      ("every one of those has a named owner and a slot in the Phase 4 sequence "
       "on slide 21. An unnamed gap becomes a production surprise.",
       {"color": INK_SOFT})]],
    size=12, spacing=1.15)
footer(s, num())

# =============================================================== 7 · TWO MODES
s = slide()
y = header(s, "The two purchase modes",
           "Parallel intents that converge on one checkout", accent=BERRY)
rows = [["", "One-time", "Subscription"],
        ["Customer intent", "Try it, or gift it. No commitment.",
         "Handle my mornings. Recurring revenue for us."],
        ["Entry point", "Juice carousel, combos, build-your-own box",
         "Plans grid on the home page, or /subscribe"],
        ["Configuration", "Cart: pick bottles and quantities",
         "Plan card: weekly | monthly, then start date and delivery days"],
        ["Where cadence is chosen", "n/a",
         "On the plan card, before checkout — and never asked again"],
        ["Price authority", "Sum of cart lines, each snapshotting its own price",
         "Server-side plan lookup. A client-supplied total is never trusted."],
        ["What checkout collects", "Pincode, delivery date + slot, address, payment",
         "Pincode, first delivery, address, payment + recurring consent"],
        ["What it creates", "An order with itemised lines",
         "An order with a plan slug and no lines — see the gap on slide 20"]]
table(s, MX, y, CW, rows, [2.3, 4.6, 5.0], row_h=0.52, head_h=0.36, size=11.5)
box(s, MX, E(y + Inches(4.28)), CW, Inches(0.72), fill=CREAM, line=LINE)
box(s, MX, E(y + Inches(4.28)), Inches(0.05), Inches(0.72), fill=BERRY)
txt(s, E(MX + Inches(0.26)), E(y + Inches(4.44)), E(CW - Inches(0.55)),
    Inches(0.5),
    [[("The design constraint:  ", {"bold": True, "color": INK}),
      ("both modes end in one checkout component running in two modes — not two "
       "checkouts. One place to get validation, serviceability and payment "
       "right, and one place where subscription-only fields must not leak into "
       "the one-time path.", {"color": INK_SOFT})]],
    size=12, spacing=1.16)
footer(s, num())

# =============================================================== 8 · JOURNEY
s = slide()
y = header(s, "End-to-end journey", "Landing to fulfilment, both paths")
j = """                            LAND ON HOME   ( no gate, no login )
                  Browse hero, seasonal, juices, combos, plans
                  Header: "Deliver to: [ pincode ]"  -- optional, non-blocking
                                        |
            +---------------------------+---------------------------+
            v                                                       v
      ONE-TIME PATH                                        SUBSCRIPTION PATH
  try / gift / no commitment                            weekly or monthly, recurring
            |                                                       |
 Pick juices, a combo, or build a box             Open Plans -> toggle Weekly | Monthly
            |                                                       |
 Add to cart -> cart drawer  ( persists )         Pick a plan -> start date + days
            |                                                       |
            +---------------------------+---------------------------+
                                        v
                            CHECKOUT   ( delivery-first )
                  1. Pincode          -> HARD serviceability check
                  2. Delivery date + slot
                  3. Address + contact
                  4. Pay              -> mock gateway, real decline path
                                        |
                     charge succeeds    |    charge declines
                                        |         `-->  nothing is written
                                        v
                  ORDER PERSISTED  +  customer profile saved
                                        |
            +---------------------------+---------------------------+
            v                                                       v
     CUSTOMER SIDE                                            OPS SIDE
  Confirmation page                                  Order appears in admin
  Account -> history + timeline                      pending -> confirmed -> pressing
  /track with order ID + email                          -> out-for-delivery -> delivered
                                                     Dashboard KPIs, low-stock alerts
                                                     ( recipe -> materials: modelled,
                                                       not yet deducted )"""
mono(s, MX, y, CW, Inches(5.42), j, size=8.7)
footer(s, num())

# =============================================================== 9 · DECISION 1
s = slide()
y = header(s, "Key decision 1 of 3",
           "Hybrid serviceability: soft while browsing, hard at checkout",
           accent=CARROT)
rows = [["Where", "Behaviour"],
        ["Home, catalogue, product pages", "No check at all. Fully browsable by anyone."],
        ["Header pincode control",
         "Optional and persisted. Shows a soft banner when out of area."],
        ["Add to cart", "Allowed regardless of pincode."],
        ["Checkout / start subscription",
         [("HARD check. ", {"bold": True, "color": APPLE}),
          ("Enforced in the server action, not just the form.", {"color": INK})]],
        ["Source of truth",
         "One function, used by the header control, the form and the action alike."]]
table(s, MX, y, HALF, rows, [2.4, 4.2], row_h=0.58, head_h=0.34, size=11,
      header_fill=CARROT)

txt(s, X2, y, HALF, Inches(0.24), "WHY NOT A PINCODE GATE", size=10,
    color=CARROT, bold=True, font=H_FONT)
txt(s, X2, E(y + Inches(0.26)), HALF, Inches(1.1),
    "A gate is the obvious answer and the wrong one: it kills the first "
    "impression and hides the entire catalogue from search. But juice is "
    "perishable and delivery-day dependent, so we genuinely can't take money "
    "for an address we can't reach. Hybrid is the only option that keeps both.",
    size=12, color=INK_SOFT, spacing=1.18)
box(s, X2, E(y + Inches(1.52)), HALF, Inches(1.64),
    fill=RGBColor(0xFD, 0xF4, 0xE8), line=CARROT)
txt(s, E(X2 + Inches(0.20)), E(y + Inches(1.68)), E(HALF - Inches(0.4)),
    Inches(0.24), "⚠  OPEN ISSUE — NEEDS A DECISION", size=10, color=APPLE,
    bold=True, font=H_FONT)
txt(s, E(X2 + Inches(0.20)), E(y + Inches(1.94)), E(HALF - Inches(0.4)),
    Inches(1.0),
    [[("The allow-list is currently switched off. ", {"bold": True, "color": INK}),
      ("A flag accepts every valid 6-digit pincode, so the real list — 10 "
       "pincodes across Mumbai, Bengaluru, Hyderabad and Kochi — the soft "
       "banner, and the whole hybrid contract are never exercised in practice. "
       "Before we flip it back, we have to decide what actually happens to an "
       "out-of-area visitor.", {"color": INK_SOFT})]],
    size=11.5, spacing=1.16)
txt(s, X2, E(y + Inches(3.20)), HALF, Inches(0.24), "THE SUB-DECISION",
    size=10, color=CARROT, bold=True, font=H_FONT)
bullets(s, X2, E(y + Inches(3.46)), HALF, [
    ("Option A", "capture email and pincode into a waitlist, then show the banner. Marketing gets a demand map by area."),
    ("Option B", "soft-warn only and let them order anyway. Simpler, and we eat the failed deliveries."),
], size=11.5, gap=0.28, dot=CARROT)
footer(s, num())

# =============================================================== 10 · DECISION 2
s = slide()
y = header(s, "Key decision 2 of 3",
           "Checkout is delivery-first, because the product is perishable",
           accent=CARROT)
cw4 = E((CW - Inches(0.36)) / 4)
steps = [("1", "Pincode", "The one hard serviceability check. Fails closed.", CARROT),
         ("2", "Date + slot", "A press-and-deliver lead time of one day; three morning slots.", GOLD),
         ("3", "Address", "Prefilled from the saved profile on repeat orders.", LEAF_SOFT),
         ("4", "Payment", "Mocked gateway, with a real decline path and a real wait.", BERRY)]
for i, (n, t, b, col) in enumerate(steps):
    x = E(MX + i * (cw4 + Inches(0.12)))
    box(s, x, y, cw4, Inches(1.78), fill=WHITE, line=LINE)
    box(s, x, y, cw4, Inches(0.05), fill=col)
    txt(s, E(x + Inches(0.20)), E(y + Inches(0.20)), Inches(0.5), Inches(0.4),
        n, size=24, color=col, bold=True, font=H_FONT)
    txt(s, E(x + Inches(0.20)), E(y + Inches(0.68)), E(cw4 - Inches(0.4)),
        Inches(0.3), t, size=13.5, color=INK, bold=True, font=H_FONT)
    txt(s, E(x + Inches(0.20)), E(y + Inches(0.98)), E(cw4 - Inches(0.4)),
        Inches(0.7), b, size=11, color=INK_SOFT, spacing=1.14)

txt(s, MX, E(y + Inches(2.06)), CW, Inches(0.24), "WHAT THIS BUYS US",
    size=10, color=CARROT, bold=True, font=H_FONT)
bullets(s, MX, E(y + Inches(2.32)), HALF, [
    ("Nobody is told “no” after paying.", "Serviceability and date both resolve before the charge."),
    ("A decline leaves nothing behind.", "Payment runs before the order is written, so there is no half-order to reconcile later."),
    ("Ops gets a press date, not an estimate.", "The delivery date is a committed field on the order rather than a computed guess."),
], size=12, gap=0.30)
txt(s, X2, E(y + Inches(2.06)), HALF, Inches(0.24), "WHAT'S STILL OPEN",
    size=10, color=APPLE, bold=True, font=H_FONT)
bullets(s, X2, E(y + Inches(2.32)), HALF, [
    ("The real cutoff rule.", "Is it “order before 8pm → next day”? How far ahead can someone book? Which days don't we deliver at all?"),
    ("Three sources currently disagree.", "The form pins the date to today, the lead-time helper says tomorrow onwards, and plan config starts tomorrow. That's a live bug, not a nuance — and the cancellation window depends on the same rule."),
], size=12, gap=0.30, dot=APPLE)
footer(s, num())

# =============================================================== 11 · DECISION 3
s = slide()
y = header(s, "Key decision 3 of 3",
           "A bill of materials is what makes the ops side worth building",
           accent=CARROT)
d = """  STOREFRONT                        SHARED DOMAIN                       OPS CONSOLE

  Customer buys            Product  "Summer Crush"  Rs.180            Materials inventory
  2 x Summer Crush   ---->    recipe:                       ---->     watermelon   stock / reorder
                                watermelon   180 g                    strawberry   stock / reorder
  The order line                  strawberry    60 g                  lime         stock / reorder
  snapshots the price             lime          15 g                  mint         stock / reorder
  at time of sale                 mint           5 g
                                                                      Low-stock alert fires when
  15 products                 26 raw materials                        stock <= reorder level
                              15 fruit / 6 veg / 5 other"""
mono(s, MX, y, CW, Inches(2.20), d, size=9.8)

txt(s, MX, E(y + Inches(2.44)), HALF, Inches(0.24), "WHY IT MATTERS", size=10,
    color=LEAF, bold=True, font=H_FONT)
bullets(s, MX, E(y + Inches(2.70)), HALF, [
    ("It closes the loop.", "Without it, the ops app is a list of orders. With it, a sale connects to a purchase order for produce."),
    ("It's why this is one domain.", "Both apps read the same product definitions, so a recipe change can never mean two different things."),
    ("It's modelled now, deliberately.", "Getting the shape right before there's data to migrate is far cheaper than retrofitting it."),
], size=11.5, gap=0.28)
txt(s, X2, E(y + Inches(2.44)), HALF, Inches(0.24), "THE HONEST CAVEAT",
    size=10, color=APPLE, bold=True, font=H_FONT)
bullets(s, X2, E(y + Inches(2.70)), HALF, [
    ("Nothing decrements stock yet.", "Recipes exist and are displayed; low-stock alerts run purely off seed values. The loop is drawn, not wired."),
    ("The fix is a ledger, not a counter.", "Stock should become a sum over append-only movements — order, restock, wastage, correction — rather than a number we mutate."),
    ("Subscriptions can't compute it at all", "until we decide whether subscribers pick their own bottles. See slide 16."),
], size=11.5, gap=0.28, dot=APPLE)
footer(s, num())

# =============================================================== 12 · SYSTEM
s = slide()
y = header(s, "For senior architects",
           "System shape: two apps, five packages, one swappable seam",
           accent=LEAF_SOFT)
a = """ +------------------------+          +--------------------------+
 |   storefront  :3000    |          |      admin  :3001        |
 |  RSC + server actions  |          |   RSC + server actions   |
 +-----------+------------+          +------------+-------------+
             |                                    |
             +-----------------+------------------+
                               v
          +----------------------------------------+
          |  @glowroot/data                        |  <-- THE SEAM
          |  every accessor async, every return     |      mock today,
          |  value a @glowroot/types shape          |      Postgres later
          +----+-----------------------------+-----+
               |                             |
     +---------v--------+        +------------v--------------+
     | in-memory seeds  |        | file-backed stores        |
     | products         |        | orders / users / sessions |
     | materials        |        +------------+--------------+
     | combos, plans    |                     |
     | produce          |          JSON files at the repo root
     +------------------+"""
mono(s, MX, y, Inches(6.5), Inches(3.30), a, size=9.4)

wr = E(CW - Inches(6.80))
xr = E(MX + Inches(6.80))
txt(s, xr, y, wr, Inches(0.24), "THE DELIBERATE CHOICE", size=10, color=LEAF,
    bold=True, font=H_FONT)
txt(s, xr, E(y + Inches(0.26)), wr, Inches(0.95),
    "There is no separate backend service. No Express app, no API process, no "
    "database. Server-side work happens in Next.js server actions, and the data "
    "layer is simply a package both apps import.",
    size=12, color=INK_SOFT, spacing=1.18)
box(s, xr, E(y + Inches(1.34)), wr, Inches(1.96), fill=CREAM, line=LINE)
txt(s, E(xr + Inches(0.20)), E(y + Inches(1.50)), E(wr - Inches(0.4)),
    Inches(0.24), "WHY THAT HOLDS UP", size=10, color=LEAF, bold=True,
    font=H_FONT)
txt(s, E(xr + Inches(0.20)), E(y + Inches(1.76)), E(wr - Inches(0.4)),
    Inches(1.4),
    "Because the accessor signatures are already the shape a real backend would "
    "expose — all async, all returning domain types. Swapping the bodies for "
    "database queries changes nothing above them. A standalone API only earns "
    "its keep when there's a non-Next consumer, and even then route handlers "
    "cover gateway webhooks fine.",
    size=11.5, color=INK_SOFT, spacing=1.16)

txt(s, MX, E(y + Inches(3.44)), CW, Inches(0.24), "WHAT I WANT YOUR EYES ON",
    size=10, color=APPLE, bold=True, font=H_FONT)
rows = [["Issue", "Where it bites", "My proposed answer"],
        ["Read-modify-write JSON with no locking",
         "Two concurrent checkouts and one is lost; a crash mid-write truncates the file",
         "Postgres behind the same accessors — Phase 4 step 3"],
        ["The data package exports fs and crypto from its root",
         "One careless import from a client component breaks the build. Worse with a DB client in there.",
         "Split into a client-safe catalogue entry and a server-only entry"],
        ["No idempotency on order placement",
         "Double-submit means two charges and two orders; the guard is client-side only",
         "A unique request ID per attempt, inserted under a unique index"]]
table(s, MX, E(y + Inches(3.70)), CW, rows, [3.0, 4.9, 4.1], row_h=0.46,
      head_h=0.32, size=10.5, header_fill=LEAF_SOFT)
footer(s, num())

# =============================================================== 13 · DOMAIN
s = slide()
y = header(s, "For senior architects",
           "Domain model: the invariants I want defended", accent=LEAF_SOFT)
txt(s, MX, y, HALF, Inches(0.24), "THE SIX ENTITIES", size=10, color=LEAF,
    bold=True, font=H_FONT)
bullets(s, MX, E(y + Inches(0.26)), HALF, [
    ("Product", "slug-keyed, and carries a recipe of raw materials"),
    ("Material", "raw produce, with a stock level and a reorder threshold"),
    ("Combo / Bundle / Plan", "curated sets, and the three subscription tiers"),
    ("Order", "customer, items, type, status, total, delivery date and slot, optional payment, append-only history"),
    ("Payment", "method, status, gateway reference, masked instrument hint"),
    ("User / Session", "email-keyed user; opaque server session with a hashed token"),
], size=11.5, gap=0.26)
txt(s, MX, E(y + Inches(2.46)), HALF, Inches(0.24),
    "FOUR INVARIANTS ALREADY IN PLACE", size=10, color=LEAF_SOFT, bold=True,
    font=H_FONT)
bullets(s, MX, E(y + Inches(2.72)), HALF, [
    ("Order lines snapshot their price.", "A later catalogue price change cannot rewrite history."),
    ("Order history is append-only.", "It's the audit trail and the customer-facing timeline in one object."),
    ("Only a masked instrument hint is stored.", "Never a full card number or UPI handle."),
    ("The password hash never crosses the boundary.", "A stripped public shape is the only user object a client can ever see."),
], size=11.5, gap=0.26, dot=LEAF_SOFT)

txt(s, X2, y, HALF, Inches(0.24), "THE MODELLING DEBT, WORST FIRST", size=10,
    color=APPLE, bold=True, font=H_FONT)
rows = [["Debt", "Consequence"],
        ["Email is the primary key",
         "There is no stable user ID. Change an email and the order history detaches; every future foreign key has nothing to point at."],
        ["A subscription is not an entity",
         "Cadence, start date and days are used to price and display, then dropped. Nothing generates delivery two; there's nothing to pause."],
        ["Stock is a mutable number",
         "Not auditable and not idempotency-friendly. Should be a cached sum over an append-only movement ledger."],
        ["Order IDs are four random digits",
         "A 9,000-value space with no uniqueness check on insert. Collisions are a matter of when, not if."],
        ["Status transitions are unconstrained",
         "Any status can be set from any other. The lifecycle graph should be enforced server-side."]]
table(s, X2, E(y + Inches(0.26)), HALF, rows, [2.1, 4.3], row_h=0.72,
      head_h=0.32, size=10.5, header_fill=APPLE)
txt(s, X2, E(y + Inches(4.28)), HALF, Inches(0.5),
    [[("The one I most want fixed before there's real data:  ",
       {"bold": True, "color": INK}),
      ("the missing user ID. Cheap now, expensive after launch.",
       {"color": INK_SOFT})]],
    size=11.5, spacing=1.15)
footer(s, num())

# =============================================================== 14 · IDENTITY
s = slide()
y = header(s, "For architects & security",
           "Identity: what's real, and the one hole left", accent=LEAF_SOFT)
third = E((CW - Inches(0.36)) / 3)
card(s, MX, y, third, Inches(1.98), "Passwords are real",
     "Salted PBKDF2-SHA256 at 120,000 iterations with a 32-byte key, stored as "
     "salt plus digest. Plaintext never touches disk, and verification is "
     "constant-time with a length check first.",
     accent=LEAF_SOFT, kicker="Done")
card(s, E(MX + third + Inches(0.18)), y, third, Inches(1.98),
     "Sessions are server-side",
     "Sign-in mints 32 random bytes. The hex goes to the browser in an "
     "HTTP-only, same-site cookie; only its SHA-256 digest is stored. Opaque, "
     "so it carries no claims — and revocable, which a signed token wouldn't be.",
     accent=LEAF_SOFT, kicker="Done")
card(s, E(MX + 2 * (third + Inches(0.18))), y, third, Inches(1.98),
     "The admin app has no auth",
     "Anyone who can reach the admin port can change any order's status. No "
     "login, no role check, and no record of who did what. This is the "
     "highest-priority item on the roadmap.",
     accent=APPLE, kicker="Open — blocking")

wl = Inches(6.4)
txt(s, MX, E(y + Inches(2.24)), wl, Inches(0.24),
    "THE GUEST-TO-REGISTERED PATH — THE BIT WORTH UNDERSTANDING", size=10,
    color=LEAF, bold=True, font=H_FONT)
rows = [["State of an email", "User record", "Has credentials", "Can check out"],
        ["Unknown", "none", "—", "no"],
        ["Guest — ordered, never registered", "exists", "no", "no, today"],
        ["Registered", "exists", "yes", "yes"]]
table(s, MX, E(y + Inches(2.50)), wl, rows, [3.0, 1.3, 1.5, 1.4], row_h=0.34,
      head_h=0.32, size=10.5, header_fill=LEAF)
txt(s, MX, E(y + Inches(4.02)), wl, Inches(0.80),
    "Checking out without an account creates a passwordless record. Signing up "
    "later with the same email fills in the credentials rather than duplicating "
    "— so the customer's entire order history is there the moment they register.",
    size=11.5, color=INK_SOFT, spacing=1.16)

xr = E(MX + Inches(6.70))
wr = E(CW - Inches(6.70))
txt(s, xr, E(y + Inches(2.24)), wr, Inches(0.24),
    "TWO TRADE-OFFS I WANT SIGNED OFF, NOT INHERITED", size=10, color=APPLE,
    bold=True, font=H_FONT)
bullets(s, xr, E(y + Inches(2.50)), wr, [
    ("Email enumeration is accepted.", "Sign-in reveals whether an email is registered, so the UI can offer sign-up instead of a dead end. More informative than production auth usually is. The UX value is real; the mitigation is rate limiting, which we don't have yet."),
    ("Guest checkout contradicts itself.", "The UI offers “continue as guest” and guest records get created, but placing an order rejects an unregistered email. Pick one. I'd allow real guest checkout — the machinery exists and it converts better."),
    ("Also missing entirely:", "any password reset path, and rate limiting on sign-in."),
], size=11, gap=0.28, dot=APPLE)
footer(s, num())

# =============================================================== 15 · DESIGN SYS
s = slide()
y = header(s, "For UI / UX designers",
           "The design system is inherited, not invented", accent=BERRY)
txt(s, MX, y, CW, Inches(0.40),
    "The original hand-built prototype already had a coherent token system. It "
    "was lifted straight into a shared Tailwind preset, so both apps speak one "
    "design language from one file — the look didn't change, we just gained "
    "reusable parts.",
    size=12.5, color=INK_SOFT, spacing=1.18)
txt(s, MX, E(y + Inches(0.56)), CW, Inches(0.24), "COLOUR TOKENS", size=10,
    color=BERRY, bold=True, font=H_FONT)
swatches = [("leaf", LEAF, "#2F5D34", WHITE), ("leaf.deep", LEAF_DEEP, "#193D20", WHITE),
            ("leaf.soft", LEAF_SOFT, "#5A8A4A", WHITE), ("carrot", CARROT, "#D9772E", WHITE),
            ("berry", BERRY, "#7A4A8F", WHITE), ("apple", APPLE, "#D4533A", WHITE),
            ("gold", GOLD, "#E8B23F", INK), ("cream", CREAM, "#F6F0E2", INK),
            ("paper", PAPER, "#FBF8F0", INK), ("ink", INK, "#243018", WHITE),
            ("ink.soft", INK_SOFT, "#5B6450", WHITE), ("line", LINE, "#E0D6BF", INK)]
sw = E((CW - Inches(0.10) * 11) / 12)
for i, (name, col, hexv, tcol) in enumerate(swatches):
    x = E(MX + i * (sw + Inches(0.10)))
    box(s, x, E(y + Inches(0.82)), sw, Inches(0.80), fill=col,
        line=LINE if col in (CREAM, PAPER, LINE) else None)
    txt(s, x, E(y + Inches(1.00)), sw, Inches(0.24), name, size=9, color=tcol,
        bold=True, font=H_FONT, align=PP_ALIGN.CENTER)
    txt(s, x, E(y + Inches(1.24)), sw, Inches(0.22), hexv, size=8, color=tcol,
        font=M_FONT, align=PP_ALIGN.CENTER)

txt(s, MX, E(y + Inches(1.86)), HALF, Inches(0.24),
    "TYPE & COMPONENT INVENTORY", size=10, color=BERRY, bold=True, font=H_FONT)
bullets(s, MX, E(y + Inches(2.12)), HALF, [
    ("Three type families", "a display face for headings, a body face, and a script accent — all set in the shared preset"),
    ("6 shared components", "button, badge, card, eyebrow, product card, stats band — imported by both apps"),
    ("33 storefront components", "10 are direct ports of prototype sections; the rest are the journey the prototype only hinted at"),
    ("3 admin components", "nav, status badge, status changer — intentionally thin"),
    ("15 storefront routes, 5 admin routes", "including build-your-own box, account order history, and public tracking"),
], size=11, gap=0.26, dot=BERRY)
txt(s, X2, E(y + Inches(1.86)), HALF, Inches(0.24),
    "WHAT MUST BE INTERACTIVE ON THE CLIENT", size=10, color=BERRY, bold=True,
    font=H_FONT)
bullets(s, X2, E(y + Inches(2.12)), HALF, [
    ("Juice carousel", "and the seasonal produce rings"),
    ("Weekly / monthly billing toggle", "cadence is chosen here, on the plan card, and never asked again"),
    ("Cart drawer and fly-to-cart", "cart and pincode both persist across reloads"),
    ("Build-your-own box builder", "two different fills of the same box stay separate cart lines"),
    ("Checkout form", "inline validation, gated on touch and submit rather than firing on every keystroke"),
    ("Mobile drawer and shop mega-menu", "two columns: seasonal, and juices + bundles"),
], size=11, gap=0.26, dot=BERRY)
footer(s, num())

# =============================================================== 16 · UX OPEN
s = slide()
y = header(s, "For UI / UX designers",
           "The states we owe, and the calls I need from you", accent=BERRY)
txt(s, MX, y, HALF, Inches(0.24), "STATES ALREADY HANDLED", size=10,
    color=LEAF_SOFT, bold=True, font=H_FONT)
bullets(s, MX, E(y + Inches(0.26)), HALF, [
    ("Empty cart", "in the drawer"),
    ("Unknown order", "on both the confirmation and tracking paths"),
    ("Invalid pincode and inline field errors", "at checkout"),
    ("Payment decline", "the mock gateway has a real failure path, deliberately triggerable from the UI"),
    ("Order timeline", "renders even for seed orders that predate history tracking"),
], size=11, gap=0.26, dot=LEAF_SOFT)
txt(s, MX, E(y + Inches(2.10)), HALF, Inches(0.24), "STATES STILL MISSING",
    size=10, color=APPLE, bold=True, font=H_FONT)
bullets(s, MX, E(y + Inches(2.36)), HALF, [
    ("Out-of-area", "the contract promises “join the waitlist”; there is no waitlist, so there's nothing to design against yet"),
    ("Logged-out header", "the account control renders nothing when signed out, so the login page is reachable only by typing the URL"),
    ("Subscription management", "no pause, skip, swap or cancel screens — blocked on the model decision"),
    ("Notifications", "no confirmation email or delivery-day reminder to design"),
], size=11, gap=0.26, dot=APPLE)

txt(s, X2, y, HALF, Inches(0.24), "FIVE CALLS I NEED FROM DESIGN", size=10,
    color=BERRY, bold=True, font=H_FONT)
rows = [["Question", "My leaning"],
        ["Should a logged-out visitor be able to find the login page?",
         "Yes — add the entry point. If hiding it was deliberate, let's write that down so it stops reading as a bug."],
        ["“Add to cart” or “add to order”?",
         "Decide once and apply everywhere: button, drawer title, and the noun in all copy."],
        ["Cart drawer only, or a full /cart page too?",
         "Drawer-first is fine, but a cart can't be deep-linked today and there's no full-page review step."],
        ["Do subscribers pick their own bottles?",
         "The contract says mix and match, the form doesn't collect it, and the confirmation says we choose. This also decides whether ops can forecast produce for subscribers at all."],
        ["Should the soft banner escalate?",
         "Today someone out of area can browse, fill the whole checkout, and only then be blocked. That's the intended hybrid — worth confirming it still feels right."]]
table(s, X2, E(y + Inches(0.26)), HALF, rows, [2.4, 4.0], row_h=0.80,
      head_h=0.32, size=10.5, header_fill=BERRY)
footer(s, num())

# =============================================================== 17 · MARKETING
s = slide()
y = header(s, "For marketing",
           "What the site does for acquisition — and the one leak", accent=GOLD)
txt(s, MX, y, HALF, Inches(0.24), "WHAT THE ARCHITECTURE GIVES YOU", size=10,
    color=CARROT, bold=True, font=H_FONT)
bullets(s, MX, E(y + Inches(0.26)), HALF, [
    ("A fully crawlable catalogue.", "No gate and no login wall. Every product, combo and seasonal page is server-rendered and indexable."),
    ("A balanced home page.", "One-time and subscription get equal weight and cross-link, so a gifter can become a subscriber without starting over."),
    ("Seasonality as data, not copy.", "Peak-season produce is modelled, so “what's in season now” can drive a campaign without a code change."),
    ("Pressed-to-order as the story.", "Next-day delivery and a real press date are structural facts we can promise, not claims we have to soften."),
], size=11, gap=0.26)
txt(s, MX, E(y + Inches(2.30)), HALF, Inches(0.24),
    "MESSAGING THE PRODUCT CAN BACK UP", size=10, color=CARROT, bold=True,
    font=H_FONT)
txt(s, MX, E(y + Inches(2.56)), HALF, Inches(0.8),
    "Pressed the night before  ·  Delivered before breakfast  ·  Peak-season "
    "produce  ·  Three plans, from a starter to a family reset  ·  Build your "
    "own box  ·  Gift it without a subscription",
    size=12, color=INK_SOFT, spacing=1.3)
box(s, MX, E(y + Inches(3.52)), HALF, Inches(1.32),
    fill=RGBColor(0xFD, 0xF4, 0xE8), line=CARROT)
txt(s, E(MX + Inches(0.20)), E(y + Inches(3.68)), E(HALF - Inches(0.4)),
    Inches(0.24), "⚠  THE LEAK", size=10, color=APPLE, bold=True, font=H_FONT)
txt(s, E(MX + Inches(0.20)), E(y + Inches(3.94)), E(HALF - Inches(0.4)),
    Inches(0.72),
    "We promise out-of-area visitors a waitlist and capture nothing. Everyone "
    "who lands, wants it, and can't buy is lost silently — and we lose the "
    "demand-by-pincode map that would tell us which city to open next.",
    size=11.5, color=INK_SOFT, spacing=1.16)

txt(s, X2, y, HALF, Inches(0.24), "THE FUNNEL, AND WHERE IT BREAKS TODAY",
    size=10, color=CARROT, bold=True, font=H_FONT)
rows = [["Stage", "Status", "Gap"],
        ["Land & browse", status_row("done"), "None — this part is solid"],
        ["Consider — detail, seasonal, combos", status_row("done"), "None"],
        ["Add to cart", status_row("done"), "None"],
        ["Check out", status_row("partial"), "Requires an account; guest checkout is half-built"],
        ["Confirm", status_row("partial"), "On-screen only — no confirmation email"],
        ["Retain / repeat", status_row("partial"), "Order history exists; no reorder button"],
        ["Subscribe & stay", status_row("missing"), "No recurring delivery, no pause or skip"],
        ["Out of area", status_row("missing"), "No waitlist capture at all"]]
table(s, X2, E(y + Inches(0.26)), HALF, rows, [2.5, 1.15, 2.9], row_h=0.42,
      head_h=0.32, size=10.5, header_fill=CARROT)
txt(s, X2, E(y + Inches(4.12)), HALF, Inches(0.6),
    [[("What I need from you:  ", {"bold": True, "color": INK}),
      ("the launch cities — which decides when we switch the pincode list back "
       "on — and whether a confirmation email is launch-blocking. I think it is.",
       {"color": INK_SOFT})]],
    size=11.5, spacing=1.16)
footer(s, num())

# =============================================================== 18 · QA
s = slide()
y = header(s, "For QA", "Test strategy: the paths, the matrix, and what's mocked",
           accent=APPLE)
third = E((CW - Inches(0.36)) / 3)
txt(s, MX, y, third, Inches(0.24), "CRITICAL PATHS TO COVER", size=10,
    color=APPLE, bold=True, font=H_FONT)
bullets(s, MX, E(y + Inches(0.26)), third, [
    ("One-time order, happy path", "browse → cart → all four checkout steps → order visible in admin"),
    ("Subscription order", "plan → cadence → config → checkout. Confirm the price comes from the server, not the form."),
    ("Payment decline", "confirm nothing is written — no order, and no profile update"),
    ("Guest → registered", "order as a guest, sign up with the same email, confirm history attaches and no duplicate appears"),
    ("Public tracking", "order ID alone must fail; ID plus email must succeed"),
    ("Cart persistence", "reload mid-cart, and confirm two different box fills stay separate lines"),
], size=10.5, gap=0.24, dot=APPLE)

xb = E(MX + third + Inches(0.18))
txt(s, xb, y, third, Inches(0.24), "NEGATIVE & EDGE CASES", size=10,
    color=APPLE, bold=True, font=H_FONT)
bullets(s, xb, E(y + Inches(0.26)), third, [
    ("Call the server actions directly.", "The UI gate is a convenience; the action is the boundary. Every action must re-validate rather than trust the form."),
    ("Try to order as somebody else.", "Identity now comes off the server session, so a request naming another email must be ignored, not honoured."),
    ("Double-submit checkout.", "Expected to fail today — no idempotency key, so this should produce two charges. Log it; don't fix it in test."),
    ("Two concurrent orders.", "Against the JSON store, one will be lost. Known, and the reason the database step exists."),
    ("Empty cart, invalid pincode, missing date.", "All three must be rejected server-side, not only in the form."),
], size=10.5, gap=0.24, dot=APPLE)

xc = E(MX + 2 * (third + Inches(0.18)))
txt(s, xc, y, third, Inches(0.24), "THE STATUS MATRIX", size=10, color=APPLE,
    bold=True, font=H_FONT)
mono(s, xc, E(y + Inches(0.26)), third, Inches(1.28),
     "pending\n   |\n   v\nconfirmed -> pressing -> out-for-delivery\n"
     "                              |\n                              v\n"
     " cancelled <---------------  delivered", size=8.4)
txt(s, xc, E(y + Inches(1.66)), third, Inches(0.95),
    [[("Test that any status can be set from any status — ", {"color": INK_SOFT}),
      ("because today it can.", {"bold": True, "color": APPLE}),
      (" Transitions are unconstrained and the graph isn't enforced. Confirm "
       "each change appends to the order's history, and that the customer's "
       "timeline reflects it.", {"color": INK_SOFT})]],
    size=10.5, spacing=1.16)
box(s, xc, E(y + Inches(2.72)), third, Inches(1.92), fill=CREAM, line=LINE)
txt(s, E(xc + Inches(0.18)), E(y + Inches(2.88)), E(third - Inches(0.36)),
    Inches(0.24), "WHAT'S MOCKED — SO TEST THE SEAM", size=10, color=LEAF,
    bold=True, font=H_FONT)
txt(s, E(xc + Inches(0.18)), E(y + Inches(3.14)), E(third - Inches(0.36)),
    Inches(1.45),
    "Payment is a stand-in that validates the instrument shape, waits about a "
    "second so “contacting your bank” is a real wait, and returns a reference. "
    "It already returns the exact shape the order persists — so test against "
    "that contract, not the fake. Then the real gateway swap is a body change "
    "with your suite still green.",
    size=10.5, color=INK_SOFT, spacing=1.16)

txt(s, MX, E(y + Inches(4.52)), Inches(7.85), Inches(0.5),
    [[("Non-functional, and honest about it:  ", {"bold": True, "color": INK}),
      ("there's no CI, no automated suite and no load testing today. The first "
       "things worth automating are the one-time happy path and the decline "
       "path — between them they catch most regressions across the four write "
       "paths.", {"color": INK_SOFT})]],
    size=11.5, spacing=1.16)
footer(s, num())

# =============================================================== 19 · DEVOPS
s = slide()
y = header(s, "For DevOps", "What runs today, and what production actually needs",
           accent=TEAL)
txt(s, MX, y, HALF, Inches(0.24), "TODAY", size=10, color=TEAL, bold=True,
    font=H_FONT)
rows = [["Concern", "Current state"],
        ["Runtime", "Two Next.js processes — storefront on 3000, admin on 3001"],
        ["Build & tasks", "A Turborepo pipeline over a pnpm workspace, with caching"],
        ["Persistence", "Two JSON files at the repo root, shared by both processes"],
        ["Secrets", "None — there's nothing to hold yet"],
        ["CI / CD", "None"],
        ["Observability", "None. No logs, no metrics, no error tracking."],
        ["Environments", "Local only"]]
table(s, MX, E(y + Inches(0.26)), HALF, rows, [1.8, 4.6], row_h=0.40,
      head_h=0.32, size=10.5, header_fill=TEAL)
box(s, MX, E(y + Inches(3.56)), HALF, Inches(1.44),
    fill=RGBColor(0xFD, 0xF4, 0xE8), line=CARROT)
txt(s, E(MX + Inches(0.20)), E(y + Inches(3.72)), E(HALF - Inches(0.4)),
    Inches(0.24), "⚠  THE GOTCHA THAT WILL COST SOMEBODY AN AFTERNOON",
    size=10, color=APPLE, bold=True, font=H_FONT)
txt(s, E(MX + Inches(0.20)), E(y + Inches(3.98)), E(HALF - Inches(0.4)),
    Inches(0.95),
    "Cookies ignore port. In local development the storefront and the admin app "
    "share one cookie jar, so their sessions collide. The answer is distinct "
    "cookie names plus validating the session's role on read — never trusting "
    "the cookie name alone.",
    size=11.5, color=INK_SOFT, spacing=1.16)

txt(s, X2, y, HALF, Inches(0.24), "WHAT PHASE 4 WILL ASK YOU FOR", size=10,
    color=TEAL, bold=True, font=H_FONT)
bullets(s, X2, E(y + Inches(0.26)), HALF, [
    ("A managed Postgres per environment", "plus a migration step in the deploy. The JSON stores go away."),
    ("Two deploy targets", "one project per app; the monorepo makes this straightforward"),
    ("Secrets management", "gateway keys, the database URL, and the session configuration"),
    ("A public webhook route", "the gateway needs a callback endpoint — the first thing here that must be reachable from outside"),
    ("A cookie-domain strategy", "storefront and admin need separate cookie names and a real domain split, not just ports"),
    ("A data migration window", "the existing orders and users files need a one-shot import. Password hashes are portable, so nobody has to reset one."),
    ("Rate limiting on sign-in", "per email and per IP. An in-memory window works to start but won't hold across instances — flag that when we scale out."),
    ("Backups and a restore drill", "the first day we have data worth losing is the day Postgres lands"),
], size=10.5, gap=0.24, dot=TEAL)
footer(s, num())

# =============================================================== 20 · STATUS
s = slide()
y = header(s, "Where the build stands", "Honest status, by area")
rows = [["Area", "Status", "Detail"],
        ["Catalogue & browsing", status_row("done"),
         "Home, products, detail pages, seasonal, combos, build-your-own box — all data-driven"],
        ["Cart", status_row("done"), "Drawer, quantities, persistence, fly-to-cart"],
        ["One-time checkout", status_row("done"),
         "Four steps, inline validation, order persisted, confirmation page"],
        ["Payment", status_row("partial"),
         "Mocked by design, with a real decline path behind a swappable seam"],
        ["Serviceability", status_row("partial"),
         "Built and wired, but the allow-list is switched off by a flag"],
        ["Delivery-date rule", status_row("partial"),
         "Three sources disagree — the form says today, the helpers say tomorrow onwards. A live bug."],
        ["Subscriptions", status_row("partial"),
         "Plans, cadence and config work; one order is written and the recurrence is dropped"],
        ["Accounts & identity", status_row("done"),
         "Real passwords, server sessions, order history, timeline, public tracking"],
        ["Ops console", status_row("done"),
         "Orders list and detail, status changes, KPI dashboard, materials with low-stock alerts"],
        ["Admin auth", status_row("missing"), "None at all. Top of the roadmap."],
        ["Inventory movement", status_row("missing"),
         "Recipes are modelled and displayed; stock never changes"],
        ["Notifications", status_row("missing"), "On-screen confirmation only"]]
table(s, MX, y, CW, rows, [2.4, 1.15, 8.3], row_h=0.375, head_h=0.32, size=10.5)
txt(s, MX, E(y + Inches(4.98)), CW, Inches(0.4),
    [[("Read this as the deck's credibility test.  ", {"bold": True, "color": INK}),
      ("Nine of the twelve rows are green or amber, and the three reds are all "
       "named on the roadmap. Nothing on this slide is a surprise to me — that "
       "is the whole point of having it.", {"color": INK_SOFT})]],
    size=12, spacing=1.16)
footer(s, num())

# =============================================================== 21 · ROADMAP
s = slide()
y = header(s, "Roadmap", "Phase 4: the sequence, and why this order")
rows = [["#", "Step", "Why it goes here", "Rough size"],
        ["1", "Admin authentication and role gating",
         "The only outright security bug left. Needs no database — it works against the current stores, and the customer-side session work already proves the pattern.",
         "~1 day"],
        ["2", "Split the data package's export surface",
         "Client-safe catalogue versus server-only stores, so misuse fails at build time. Cheap, and a prerequisite for a DB client living in that package.",
         "~1 hour"],
        ["3", "Postgres behind the same accessors, plus migration scripts",
         "Fixes concurrent-write loss and the truncation risk in one move. The accessor signatures don't change, so no app code does either.",
         "~2–3 days"],
        ["4", "Idempotency key, real order sequence, one transaction",
         "Falls out naturally once there's a database. Kills the double-charge and the ID collision together.",
         "~half day"],
        ["5", "Subscription entity and inventory ledger",
         "The two genuinely new domain models. Unblocks pause/skip/cancel, recurring deliveries, and produce forecasting.",
         "~2 days"],
        ["6", "Real gateway, password reset, rate limiting",
         "Independent of each other — slot them in as needed once the foundation is there.",
         "as needed"]]
table(s, MX, y, CW, rows, [0.45, 3.4, 7.0, 1.15], row_h=0.62, head_h=0.32,
      size=10.5)
box(s, MX, E(y + Inches(4.16)), CW, Inches(0.90), fill=CREAM, line=LINE)
box(s, MX, E(y + Inches(4.16)), Inches(0.05), Inches(0.90), fill=LEAF)
txt(s, E(MX + Inches(0.26)), E(y + Inches(4.32)), E(CW - Inches(0.55)),
    Inches(0.55),
    [[("The sequencing argument:  ", {"bold": True, "color": INK}),
      ("steps 1 and 2 come before the database on purpose. Sessions change the "
       "accessor signatures, and I'd rather change them once against JSON files "
       "than twice against Postgres. Doing the security fix first also means we "
       "are never knowingly running an unauthenticated admin panel against a "
       "real database.", {"color": INK_SOFT})]],
    size=11.5, spacing=1.16)
footer(s, num())

# =============================================================== 22 · RISKS
s = slide()
y = header(s, "Risks & open decisions", "Six things that need an owner and a date",
           accent=APPLE)
rows = [["Risk / decision", "Impact if we do nothing", "Call needed from", "My recommendation"],
        ["Admin panel is unauthenticated",
         "Anyone reaching the port can alter any order. Blocks any real launch.",
         "Architecture + me",
         "Fix it first, before the database. Roles rather than a shared password, and record who changed what."],
        ["A subscription is not an entity",
         "No second delivery is ever generated, there's nothing to pause, and ops can't tell what to press for a subscriber.",
         "Me + architecture",
         "Build the real entity that generates an order per delivery. Get the shape right before there's data to migrate."],
        ["The delivery-date rule contradicts itself",
         "Subscription checkout is internally inconsistent today, and it's the rule cancellation windows depend on.",
         "Me + ops",
         "Define the cutoff — “order before 8pm → next day” — then make one helper the single source of truth."],
        ["Serviceability is switched off",
         "The hybrid contract is never exercised, and out-of-area demand is lost silently.",
         "Marketing + me",
         "Decide the launch cities and the out-of-area experience, add waitlist capture, then flip the flag back."],
        ["Guest checkout contradicts itself",
         "The UI offers it and the server rejects it. Measurable conversion loss.",
         "Design + me",
         "Allow real guest checkout, with claim-by-email on signup. The machinery already exists."],
        ["No confirmation email",
         "Customers have no record of an order once they close the tab.",
         "Marketing + DevOps",
         "Treat it as launch table stakes. Order placement is the natural hook."]]
table(s, MX, y, CW, rows, [2.4, 3.2, 1.75, 4.55], row_h=0.72, head_h=0.32,
      size=10, header_fill=APPLE)
footer(s, num())

# =============================================================== 23 · ASKS
s = slide(LEAF_DEEP)
box(s, 0, 0, W, Inches(0.10), fill=CARROT)
box(s, MX, Inches(0.58), Inches(0.055), Inches(0.52), fill=CARROT)
txt(s, E(MX + Inches(0.20)), Inches(0.56), CW, Inches(0.24),
    "WHAT I NEED FROM YOU", size=10.5, color=GOLD, bold=True, font=H_FONT)
txt(s, E(MX + Inches(0.20)), Inches(0.80), CW, Inches(0.42),
    "One ask each — and the date I need it by", size=24, color=CREAM,
    bold=True, font=H_FONT)
asks = [
    ("Senior architects", LEAF_SOFT,
     "Review the three issues on slide 12 and tell me whether “no separate "
     "service” still holds at ten times the traffic. Then rule on the missing "
     "user ID — I want that fixed before there's data to migrate."),
    ("UI / UX designers", BERRY,
     "The five calls on slide 16 — especially whether subscribers pick their "
     "own bottles, because that one decides whether ops can forecast produce "
     "at all. Plus the out-of-area and logged-out states we have nothing for."),
    ("Marketing", GOLD,
     "Launch cities, so I know when to switch the pincode allow-list back on. "
     "And a yes or no on the confirmation email being launch-blocking — I "
     "believe it is."),
    ("QA", APPLE,
     "Automate the one-time happy path and the payment decline path first. "
     "Then treat the documented gaps on slide 18 as expected failures rather "
     "than new bugs, and log them against roadmap items."),
    ("DevOps", TEAL,
     "A cost and lead time for managed Postgres plus two deploy targets, and "
     "the shape of secrets management. The cookie-port collision on slide 19 "
     "is the one thing I'd flag before you start."),
]
cy = Inches(1.58)
for role, col, ask in asks:
    box(s, MX, cy, CW, Inches(0.94), fill=DEEPCARD)
    box(s, MX, cy, Inches(0.045), Inches(0.94), fill=col)
    txt(s, E(MX + Inches(0.26)), E(cy + Inches(0.18)), Inches(2.6),
        Inches(0.3), role, size=13, color=col, bold=True, font=H_FONT)
    txt(s, E(MX + Inches(3.05)), E(cy + Inches(0.14)), E(CW - Inches(3.35)),
        Inches(0.72), ask, size=11.5,
        color=RGBColor(0xC2, 0xD1, 0xB9), spacing=1.16)
    cy = E(cy + Inches(1.00))

txt(s, MX, Inches(6.62), CW, Inches(0.3),
    "Every open question in this deck is written down in the repo docs — "
    "none of it lives only in my head.",
    size=11, color=RGBColor(0x86, 0x9A, 0x7C), font=B_FONT, italic=True)
footer(s, num(), dark=True)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(prs.slides._sldIdLst))
